# -*- coding: utf-8 -*-
"""قراءة Data.xls و 2026 NUMBER.xlsx وبناء الفهارس والبحث الاحترافي.

يعمل على أي بايثون (32/64 بت) — يعتمد xlrd (.xls قديم) و openpyxl (.xlsx).
"""
import os
import unicodedata

try:
    import xlrd
except ImportError:  # pragma: no cover
    xlrd = None
try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


def _norm_ar(s):
    """تطبيع نص عربي للبحث التشابهى: إزالة التشكيل وتوحيد الألف/الهاء/الياء."""
    if s is None:
        return ""
    s = str(s)
    s = "".join(c for c in unicodedata.normalize("NFD", s)
                if unicodedata.category(c) != "Mn")
    s = s.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    s = s.replace("ى", "ي").replace("ة", "ه")
    return s.lower().strip()


def _as_int_str(v):
    """تحويل خلية Excel (قد تكون float مثل 9621005209.0) إلى سلسلة أرقام نظيفة."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s


class DataStore:
    def __init__(self, data_xls=None, employees_xlsx=None):
        base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        self.data_xls = data_xls or os.path.join(base, "data", "Data.xls")
        self.employees_xlsx = employees_xlsx or os.path.join(
            base, "data", "2026 NUMBER.xlsx")
        self.schools = []          # [(code:str, name:str)]
        self.employees = {}        # ministry(str) -> record dict
        self.by_national = {}      # national(str) -> record dict
        self.load_errors = []

    # ---------- المدارس ----------
    def load_schools(self):
        if xlrd is None:
            self.load_errors.append("xlrd غير مثبت — لا يمكن قراءة Data.xls")
            return
        wb = xlrd.open_workbook(self.data_xls)
        sh = wb.sheet_by_index(0)
        out = []
        for r in range(1, sh.nrows):
            code = _as_int_str(sh.cell_value(r, 0))
            name = str(sh.cell_value(r, 1)).strip()
            if code and name:
                out.append((code, name))
        self.schools = out

    # ---------- الموظفين ----------
    def load_employees(self):
        if openpyxl is None:
            self.load_errors.append("openpyxl غير مثبت — لا يمكن قراءة 2026 NUMBER.xlsx")
            return
        wb = openpyxl.load_workbook(self.employees_xlsx, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        idx = {h: i for i, h in enumerate(header)}

        def g(row, key):
            i = idx.get(key)
            return row[i] if i is not None and i < len(row) else None

        for row in rows:
            if row is None:
                continue
            ministry = _as_int_str(g(row, "الرقم الوزاري"))
            national = _as_int_str(g(row, "الرقم الوطني"))
            if not ministry and not national:
                continue
            rec = {
                "ministry": ministry,
                "national": national,
                "name": str(g(row, "الاسم") or "").strip(),
                "gender": _as_int_str(g(row, "الجنس")),
                "work_center": str(g(row, "مركز العمل") or "").strip(),
                "directorates": str(g(row, "المديرية") or "").strip(),
                "qual_first": str(g(row, "المؤهل العلمي الأول") or "").strip(),
                "spec_first": str(g(row, "تخصص المؤهل الأول") or "").strip(),
                "qual_last": str(g(row, "المؤهل العلمي الأخير") or "").strip(),
                "spec_last": str(g(row, "تخصص المؤهل الأخير") or "").strip(),
            }
            # بقية الأعمدة (سنوات الخدمة...) تُحفظ باسم عمودها
            for k, i in idx.items():
                if k not in ("الرقم الوزاري", "الاسم", "الرقم الوطني", "الجنس",
                             "مركز العمل", "رمز المديرية", "المديرية",
                             "رمز المؤهل الأول", "المؤهل العلمي الأول",
                             "تخصص المؤهل الأول", "رمز المؤهل الأخير",
                             "المؤهل العلمي الأخير"):
                    rec[str(k)] = _as_int_str(row[i]) if i < len(row) else ""
            if ministry:
                self.employees[ministry] = rec
            if national:
                self.by_national[national] = rec
        wb.close()

    def load_all(self):
        self.load_schools()
        self.load_employees()
        return self

    # ---------- البحث ----------
    def find_employee(self, number):
        """إظهار بيانات الموظف تلقائيًا بمجرد إدخال الرقم الوزاري أو الوطني."""
        n = _as_int_str(number)
        return self.employees.get(n) or self.by_national.get(n)

    def search_schools(self, query, limit=25):
        """بحث احترافي: بالرمز الحرفي، ثم بالاسم الكامل، ثم بالتشابه الجزئي."""
        q = str(query or "").strip()
        if not q:
            return []
        exact = [s for s in self.schools if s[0] == q]
        if exact:
            return exact
        nq = _norm_ar(q)
        starts, contains, fuzzy = [], [], []
        for code, name in self.schools:
            nn = _norm_ar(name)
            if nn == nq:
                starts.append((code, name))
            elif nn.startswith(nq):
                starts.append((code, name))
            elif nq in nn:
                contains.append((code, name))
            elif all(tok in nn for tok in nq.split()):
                fuzzy.append((code, name))
        res = starts + contains + fuzzy
        return res[:limit]
